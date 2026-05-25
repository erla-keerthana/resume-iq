import os
import re
import shutil
import tempfile
from pathlib import Path
from urllib.parse import urlparse

import pdfplumber
import requests
from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
from fpdf import FPDF
import pytesseract


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".jpg", ".jpeg", ".png"}


def detect_file_type(filename: str | None) -> str:
    """Detect file type from the filename extension."""
    if not filename:
        raise ValueError("No filename provided. Cannot detect file type.")
    # Use only the basename to prevent path traversal
    safe_name = Path(filename).name
    ext = Path(safe_name).suffix.lower()
    if ext in ALLOWED_EXTENSIONS:
        return ext
    raise ValueError(f"Unsupported file type: {ext}")


def _sanitize_text(text: str) -> str:
    """Replace common Unicode characters with Latin-1 safe equivalents."""
    replacements = {
        "\u2018": "'", "\u2019": "'",  # smart single quotes
        "\u201c": '"', "\u201d": '"',  # smart double quotes
        "\u2013": "-", "\u2014": "-",  # en/em dashes
        "\u2022": "*",                  # bullet
        "\u2026": "...",               # ellipsis
        "\u00a0": " ",                 # non-breaking space
    }
    for orig, repl in replacements.items():
        text = text.replace(orig, repl)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _txt_to_pdf(file_bytes: bytes, output_path: str) -> str:
    """Convert plain text to PDF."""
    text = file_bytes.decode("utf-8", errors="ignore")
    text = _sanitize_text(text)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in text.split("\n"):
        pdf.cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(output_path)
    return output_path


def _docx_to_pdf(file_bytes: bytes, output_path: str) -> str:
    """Convert DOCX to PDF."""
    # Write bytes to a temp docx first
    tmp_docx = output_path + ".docx"
    with open(tmp_docx, "wb") as f:
        f.write(file_bytes)

    doc = Document(tmp_docx)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for para in doc.paragraphs:
        if para.text.strip():
            safe_text = _sanitize_text(para.text)
            pdf.cell(0, 7, safe_text, new_x="LMARGIN", new_y="NEXT")
    # Extract table content (common in resumes)
    for table in doc.tables:
        for row in table.rows:
            cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells_text:
                safe_text = _sanitize_text("  |  ".join(cells_text))
                pdf.cell(0, 7, safe_text, new_x="LMARGIN", new_y="NEXT")
    pdf.output(output_path)
    os.remove(tmp_docx)
    return output_path


def _image_to_pdf(file_bytes: bytes, output_path: str) -> str:
    """Convert image (JPG/PNG) to PDF using OCR."""
    tmp_img = output_path + ".img"
    with open(tmp_img, "wb") as f:
        f.write(file_bytes)

    image = Image.open(tmp_img)
    # OCR the image to get text
    try:
        text = pytesseract.image_to_string(image)
    except EnvironmentError:
        raise ValueError(
            "OCR failed: tesseract is not installed on this system. "
            "Install it with: sudo apt-get install tesseract-ocr (Linux) "
            "or brew install tesseract (macOS)"
        )
    text = _sanitize_text(text)

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in text.split("\n"):
        pdf.cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(output_path)
    os.remove(tmp_img)
    return output_path


def _extract_text_from_pdf(pdf_path: str) -> str:
    """Extract text from a PDF file using pdfplumber."""
    text = ""
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text.strip()


def process_file(file_bytes: bytes, filename: str) -> str:
    """
    Single function for all file processing:
    1. Detect file type
    2. Convert to PDF
    3. Extract text from PDF
    Returns the extracted text.
    """
    ext = detect_file_type(filename)

    # Basic content validation: PDFs must start with %PDF
    if ext == ".pdf" and not file_bytes[:5].startswith(b"%PDF"):
        raise ValueError("File does not appear to be a valid PDF")

    # Create a temp directory for intermediate files
    tmp_dir = tempfile.mkdtemp()
    pdf_path = os.path.join(tmp_dir, "converted.pdf")

    try:
        if ext == ".pdf":
            # Already a PDF, just write it
            with open(pdf_path, "wb") as f:
                f.write(file_bytes)
        elif ext == ".txt":
            _txt_to_pdf(file_bytes, pdf_path)
        elif ext == ".docx":
            _docx_to_pdf(file_bytes, pdf_path)
        elif ext in {".jpg", ".jpeg", ".png"}:
            _image_to_pdf(file_bytes, pdf_path)

        # Extract text from the PDF
        text = _extract_text_from_pdf(pdf_path)
        return text
    finally:
        # Cleanup entire temp directory and all intermediate files
        shutil.rmtree(tmp_dir, ignore_errors=True)


# ---------------------------------------------------------------------------
# LinkedIn / GitHub Profile Fetching
# ---------------------------------------------------------------------------

_PROFILE_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml",
    "Accept-Language": "en-US,en;q=0.9",
}


def fetch_profile_text(url: str) -> str:
    """Fetch public profile page and extract readable text."""
    parsed = urlparse(url)
    host = parsed.hostname or ""

    if not parsed.scheme:
        url = "https://" + url
        parsed = urlparse(url)
        host = parsed.hostname or ""

    is_linkedin = "linkedin.com" in host
    is_github = "github.com" in host

    if not is_linkedin and not is_github:
        raise ValueError("Only LinkedIn and GitHub profile URLs are supported.")

    try:
        resp = requests.get(url, headers=_PROFILE_HEADERS, timeout=15, allow_redirects=True)
        resp.raise_for_status()
    except requests.RequestException as e:
        raise ValueError(f"Could not fetch profile: {e}")

    soup = BeautifulSoup(resp.text, "html.parser")

    # Remove script/style noise
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()

    if is_github:
        return _parse_github(soup, url)
    return _parse_linkedin(soup)


def _parse_github(soup: BeautifulSoup, url: str) -> str:
    """Extract profile info from a GitHub public profile page."""
    parts = []

    # Name
    name_el = soup.select_one(".p-name, .vcard-fullname, [itemprop='name']")
    if name_el:
        parts.append(f"Name: {name_el.get_text(strip=True)}")

    # Bio
    bio_el = soup.select_one(".p-note, .user-profile-bio, [data-bio-text]")
    if bio_el:
        parts.append(f"Bio: {bio_el.get_text(strip=True)}")

    # Org / location / links
    for item in soup.select(".vcard-detail, [itemprop='worksFor'], [itemprop='homeLocation']"):
        text = item.get_text(strip=True)
        if text:
            parts.append(text)

    # Pinned repos
    pinned = soup.select(".pinned-item-list-item-content, .js-pinned-item-list-item")
    if pinned:
        parts.append("\nPinned Repositories:")
        for pin in pinned:
            repo_name = pin.select_one(".repo")
            desc = pin.select_one(".pinned-item-desc, p")
            lang = pin.select_one("[itemprop='programmingLanguage']")
            line = f"- {repo_name.get_text(strip=True)}" if repo_name else "- repo"
            if desc:
                line += f": {desc.get_text(strip=True)}"
            if lang:
                line += f" ({lang.get_text(strip=True)})"
            parts.append(line)

    # README (profile readme)
    readme = soup.select_one("article.markdown-body")
    if readme:
        parts.append(f"\nProfile README:\n{readme.get_text(separator=' ', strip=True)[:3000]}")

    if not parts:
        # Fallback: just get visible text
        parts.append(soup.get_text(separator="\n", strip=True)[:5000])

    parts.append(f"\nProfile URL: {url}")
    return "\n".join(parts)


def _parse_linkedin(soup: BeautifulSoup) -> str:
    """Extract text from a LinkedIn public profile page."""
    # LinkedIn public profiles have limited HTML; extract what we can
    parts = []

    # Try meta tags (often available on public profiles)
    title = soup.find("title")
    if title:
        parts.append(title.get_text(strip=True))

    desc = soup.find("meta", attrs={"name": "description"})
    if desc and desc.get("content"):
        parts.append(desc["content"])

    # Main content
    main = soup.find("main") or soup.find("body")
    if main:
        text = main.get_text(separator="\n", strip=True)
        # Clean up excessive whitespace
        text = re.sub(r"\n{3,}", "\n\n", text)
        parts.append(text[:5000])

    if not parts:
        raise ValueError(
            "Could not extract content from LinkedIn profile. "
            "The profile may be private or requires login."
        )

    return "\n".join(parts)
